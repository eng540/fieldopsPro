"""Excel bulk import for canonical project BOQ and project units.

The importer is intentionally project-scoped and atomic: the workbook is fully
validated before any row is persisted. This keeps Units and Master BOQ separate.
"""
from __future__ import annotations

from io import BytesIO
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.modules.iam.dependencies import get_current_user
from app.modules.projects.models import BOQItem, Project, ProjectUnit

router = APIRouter()
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000
BOQ_HEADERS = {"code": {"code", "الكود", "رمز البند"}, "category": {"category", "التصنيف", "الفئة"}, "trade": {"trade", "التخصص"}, "description": {"description", "الوصف", "وصف البند", "البيان"}, "quantity": {"quantity", "الكمية", "الكمية الإجمالية"}, "rate": {"rate", "سعر الوحدة"}, "unit_of_measure": {"unit_of_measure", "unit", "وحدة القياس"}, "sequence": {"sequence", "التسلسل", "الترتيب"}}
UNIT_HEADERS = {"name": {"name", "اسم الوحدة"}, "code": {"code", "رمز الوحدة"}, "unit_type": {"unit_type", "unit type", "نوع الوحدة"}, "floor": {"floor", "الطابق"}, "area_sqm": {"area_sqm", "area", "المساحة", "المساحة م²", "المساحة م2"}}

def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()

def _header_map(row: tuple[Any, ...], aliases: dict[str, set[str]]) -> dict[str, int]:
    normalized = {_text(v).lower(): i for i, v in enumerate(row) if _text(v)}
    result: dict[str, int] = {}
    for field, names in aliases.items():
        for name in names:
            if name.lower() in normalized:
                result[field] = normalized[name.lower()]
                break
    return result

def _required_headers(mapping: dict[str, int], required: set[str]) -> list[str]:
    return sorted(required - mapping.keys())

def _number(value: Any, row: int, field: str, *, required: bool = True) -> float:
    if value is None or _text(value) == "":
        if required:
            raise ValueError(f"الصف {row}: الحقل «{field}» مطلوب")
        return 0.0
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"الصف {row}: «{field}» يجب أن يكون رقمًا") from exc
    if number < 0:
        raise ValueError(f"الصف {row}: «{field}» لا يمكن أن يكون سالبًا")
    return number

async def _read_workbook(file: UploadFile) -> tuple[list[tuple[Any, ...]], str]:
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=422, detail="يرجى رفع ملف Excel بصيغة .xlsx")
    raw = await file.read()
    if len(raw) > MAX_FILE_BYTES:
        raise HTTPException(status_code=413, detail="حجم ملف Excel يتجاوز 10 MB")
    try:
        wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=422, detail="ملف Excel غير صالح أو تالف") from exc
    if not wb.sheetnames:
        raise HTTPException(status_code=422, detail="ملف Excel لا يحتوي على ورقة عمل")
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) <= 1:
        raise HTTPException(status_code=422, detail="ملف Excel لا يحتوي على بيانات بعد صف العناوين")
    if len(rows) - 1 > MAX_ROWS:
        raise HTTPException(status_code=422, detail=f"الحد الأقصى للاستيراد هو {MAX_ROWS} صفًا في المرة الواحدة")
    return rows, ws.title

@router.post("/{project_id}/import/boq")
async def import_boq_excel(project_id: int, file: UploadFile = File(...), db: AsyncSession = Depends(get_db), current_user: dict = Depends(get_current_user)) -> dict[str, Any]:
    org_id = current_user["org_id"]
    project = (await db.execute(select(Project).where(Project.id == project_id, Project.org_id == org_id))).scalar_one_or_none()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    rows, sheet = await _read_workbook(file)
    mapping = _header_map(rows[0], BOQ_HEADERS)
    missing = _required_headers(mapping, {"trade", "description", "quantity", "unit_of_measure"})
    if missing:
        raise HTTPException(status_code=422, detail={"message": "أعمدة BOQ المطلوبة مفقودة", "missing": missing, "sheet": sheet})
    existing_codes = set((await db.execute(select(BOQItem.code).where(BOQItem.project_id == project_id))).scalars().all())
    pending_codes: set[str] = set(); errors: list[str] = []; prepared: list[dict[str, Any]] = []
    next_code = (await db.execute(select(func.count()).select_from(BOQItem).where(BOQItem.project_id == project_id))).scalar_one() + 1
    for row_no, row in enumerate(rows[1:], start=2):
        if not any(_text(v) for v in row):
            continue
        try:
            code = _text(row[mapping["code"]]).upper() if "code" in mapping else ""
            if not code:
                while f"BOQ-{next_code:04d}" in existing_codes or f"BOQ-{next_code:04d}" in pending_codes:
                    next_code += 1
                code = f"BOQ-{next_code:04d}"; next_code += 1
            if code in existing_codes or code in pending_codes:
                raise ValueError(f"الصف {row_no}: كود BOQ «{code}» مكرر")
            trade = _text(row[mapping["trade"]]); description = _text(row[mapping["description"]]); uom = _text(row[mapping["unit_of_measure"]])
            if not trade or not description or not uom:
                raise ValueError(f"الصف {row_no}: التخصص والوصف ووحدة القياس مطلوبة")
            quantity = _number(row[mapping["quantity"]], row_no, "الكمية الإجمالية")
            rate = _number(row[mapping["rate"]], row_no, "سعر الوحدة", required=False) if "rate" in mapping else 0.0
            sequence = int(_number(row[mapping["sequence"]], row_no, "التسلسل", required=False)) if "sequence" in mapping else 0
            prepared.append({"org_id": org_id, "project_id": project_id, "code": code, "category": _text(row[mapping["category"]]) if "category" in mapping else None, "trade": trade, "description": description, "quantity": quantity, "rate": rate, "amount": quantity * rate, "unit_of_measure": uom, "sequence": sequence})
            pending_codes.add(code)
        except ValueError as exc:
            errors.append(str(exc))
    if errors:
        raise HTTPException(status_code=422, detail={"message": "تم رفض الاستيراد؛ صحح الأخطاء ثم أعد الرفع", "errors": errors[:100], "error_count": len(errors)})
    db.add_all([BOQItem(**item) for item in prepared]); await db.flush()
    return {"project_id": project_id, "imported": len(prepared), "type": "boq", "message": f"تم استيراد {len(prepared)} بند BOQ بنجاح"}

@router.post("/{project_id}/import/units")
async def import_units_excel(project_id: int, file: UploadFile = File(...), db: AsyncSession = Depends(get_db), current_user: dict = Depends(get_current_user)) -> dict[str, Any]:
    org_id = current_user["org_id"]
    project = (await db.execute(select(Project).where(Project.id == project_id, Project.org_id == org_id))).scalar_one_or_none()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    rows, sheet = await _read_workbook(file)
    mapping = _header_map(rows[0], UNIT_HEADERS)
    missing = _required_headers(mapping, {"name", "code", "unit_type"})
    if missing:
        raise HTTPException(status_code=422, detail={"message": "أعمدة الوحدات المطلوبة مفقودة", "missing": missing, "sheet": sheet})
    existing_codes = set((await db.execute(select(ProjectUnit.code).where(ProjectUnit.project_id == project_id))).scalars().all())
    pending_codes: set[str] = set(); errors: list[str] = []; prepared: list[dict[str, Any]] = []
    for row_no, row in enumerate(rows[1:], start=2):
        if not any(_text(v) for v in row):
            continue
        try:
            name = _text(row[mapping["name"]]); code = _text(row[mapping["code"]]).upper(); unit_type = _text(row[mapping["unit_type"]])
            if not name or not code or not unit_type:
                raise ValueError(f"الصف {row_no}: اسم الوحدة ورمزها ونوعها مطلوبة")
            if code in existing_codes or code in pending_codes:
                raise ValueError(f"الصف {row_no}: رمز الوحدة «{code}» مكرر")
            floor = _text(row[mapping["floor"]]) if "floor" in mapping else ""
            if floor:
                try: floor_value = int(float(floor))
                except (TypeError, ValueError) as exc: raise ValueError(f"الصف {row_no}: «الطابق» يجب أن يكون رقمًا") from exc
            else: floor_value = None
            area = _number(row[mapping["area_sqm"]], row_no, "المساحة م²", required=False) if "area_sqm" in mapping else 0.0
            prepared.append({"org_id": org_id, "project_id": project_id, "name": name, "code": code, "unit_type": unit_type, "floor": floor_value, "area_sqm": area if area else None})
            pending_codes.add(code)
        except (ValueError, TypeError) as exc:
            errors.append(str(exc))
    if errors:
        raise HTTPException(status_code=422, detail={"message": "تم رفض الاستيراد؛ صحح الأخطاء ثم أعد الرفع", "errors": errors[:100], "error_count": len(errors)})
    db.add_all([ProjectUnit(**item) for item in prepared]); project.total_units = (project.total_units or 0) + len(prepared); await db.flush()
    return {"project_id": project_id, "imported": len(prepared), "type": "units", "message": f"تم استيراد {len(prepared)} وحدة بنجاح"}
